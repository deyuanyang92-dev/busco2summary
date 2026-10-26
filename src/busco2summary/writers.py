"""Table writers (TSV / CSV / optional XLSX)."""

from __future__ import annotations

import csv
import gzip
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence

from .log import warn


def fmt(v, digits: int = 4) -> str:
    if v is None:
        return "NA"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        return f"{v:.{digits}f}".rstrip("0").rstrip(".") if v == v else "NA"
    if isinstance(v, (list, tuple)):
        return ";".join(map(str, v))
    if isinstance(v, dict):
        return ";".join(f"{k}={x}" for k, x in v.items())
    return str(v)


def write_table(path: Path, rows: Sequence[dict], fields: Optional[List[str]] = None,
                sep: str = "\t", digits: int = 4) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = []
        for r in rows:
            for k in r:
                if k not in fields:
                    fields.append(k)
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "wt", newline="") as fh:
        w = csv.writer(fh, delimiter=sep, lineterminator="\n")
        w.writerow(fields)
        for r in rows:
            w.writerow([fmt(r.get(f), digits) for f in fields])
    return path


def write_matrix(path: Path, row_ids: Sequence[str], col_ids: Sequence[str],
                 cell, corner: str = "busco_id", transpose: bool = False) -> Path:
    """cell(row_index, col_index) -> value."""
    path.parent.mkdir(parents=True, exist_ok=True)
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "wt", newline="") as fh:
        w = csv.writer(fh, delimiter="\t", lineterminator="\n")
        if not transpose:
            w.writerow([corner] + list(col_ids))
            for i, r in enumerate(row_ids):
                w.writerow([r] + [fmt(cell(i, j)) for j in range(len(col_ids))])
        else:
            w.writerow(["sample"] + list(row_ids))
            for j, c in enumerate(col_ids):
                w.writerow([c] + [fmt(cell(i, j)) for i in range(len(row_ids))])
    return path


def write_list(path: Path, items: Iterable[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(f"{x}\n" for x in items))
    return path


def write_xlsx(path: Path, sheets: Dict[str, Sequence[dict]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        warn("openpyxl not installed: skipping --xlsx")
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if not rows:
            continue
        fields = list(rows[0].keys())
        ws.append(fields)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(f) if isinstance(r.get(f), (int, float, str)) or r.get(f) is None
                       else fmt(r.get(f)) for f in fields])
        ws.freeze_panes = "B2"
    wb.save(path)
    return True
